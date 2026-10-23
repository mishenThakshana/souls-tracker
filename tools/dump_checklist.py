import openpyxl, json, sys, collections
wb=openpyxl.load_workbook(sys.argv[1])
out={}
for name,hdr in [('Base Game',5),('Shadow of the Erdtree',3)]:
    ws=wb[name]; rows=[]
    from openpyxl.utils import range_boundaries
    red=set()
    for cf in ws.conditional_formatting:
        if any(r.dxf and r.dxf.fill and r.dxf.fill.bgColor.rgb=='FFE06666' for r in cf.rules):
            for rng in str(cf.sqref).split():
                c1,r1,c2,r2=range_boundaries(rng)
                red.update(range(r1,r2+1))
    side=[]
    for i,r in enumerate(ws.iter_rows(min_row=hdr+1,values_only=True),start=hdr+1):
        item,cat,loc,notes=r[3],r[4],r[5],r[6]
        extra=[c for c in r[7:] if c not in (None,'')]
        if extra: side.append((i,extra))
        if item is None and cat is None and loc is None and notes is None: continue
        d=ws.cell(i,4)
        hdr_=bool(d.fill and d.fill.fill_type and d.fill.fgColor.rgb=='FFA2C4C9')
        rows.append(dict(row=i,done=r[2],item=item,cat=cat,loc=loc,notes=notes,header=hdr_,ng=i in red))
    out[name]=rows
    out[name+' tips']=[str(x[1][0]) for x in side if str(x[1][0]).strip() not in ('NOTES:',)]
    print(name,len(rows))
    print(' side cells:',len(side)); 
    for s in side[:15]: print('  ',s)
    c=collections.Counter(x['loc'] for x in rows); print(' locs',len(c)); print(c.most_common(400))
    c=collections.Counter(x['cat'] for x in rows); print(' cats',len(c), c.most_common())
    odd=[x for x in rows if x['item'] is None or x['done'] not in (True,False)]
    print(' odd',len(odd), odd[:20])
json.dump(out,open(sys.argv[2],'w'),default=str,indent=1)
